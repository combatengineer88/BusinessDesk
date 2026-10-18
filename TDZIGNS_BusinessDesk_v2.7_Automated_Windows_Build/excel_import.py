"""Transactional Excel imports for TDZIGNS LLC BusinessDesk.

Header-based mapping, preview, duplicate handling, and row-level validation.
Only .xlsx/.xlsm supported; formula results must be cached by Excel.
"""
from pathlib import Path
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import sqlite3

HEADERS = {
    'customers': {
        'name': ('name','customer name','customer','company','company name','full name'),
        'email': ('email','email address','e-mail'),
        'phone': ('phone','phone number','telephone','mobile'),
        'address': ('address','street address','mailing address'),
        'notes': ('notes','comments','remarks'),
    },
    'suppliers': {
        'name': ('name','supplier name','supplier','vendor','vendor name','company','company name','business name'),
        'email': ('email','email address','e-mail'),
        'phone': ('phone','phone number','telephone','mobile'),
        'address': ('address','street address','mailing address'),
    },
    'products': {
        'sku': ('sku','product sku','item number','item #','item code','product code'),
        'barcode': ('barcode','upc','ean','gtin'),
        'name': ('name','product name','item name','product','item'),
        'description': ('description','details','product description'),
        'price': ('price','selling price','sale price','retail price','unit price'),
        'cost': ('cost','unit cost','purchase cost'),
        'qty': ('qty','quantity','stock','on hand','stock quantity','quantity on hand'),
        'supplier': ('supplier','supplier name','vendor','vendor name','primary supplier'),
    }
}
REQUIRED = {'customers': ('name',), 'products': ('sku','name','price'), 'suppliers': ('name',)}
MAX_ROWS=10000

def normalized(s):
    return ' '.join(str(s or '').strip().lower().replace('_',' ').replace('-',' ').split())

def mapping(headers,kind):
    normalized_headers=[normalized(h) for h in headers]
    out={}
    for field,aliases in HEADERS[kind].items():
        found=[i for i,h in enumerate(normalized_headers) if h in aliases]
        if len(found)>1:
            raise ValueError(f"Ambiguous Excel columns for {field}. Keep only one matching column.")
        if found:out[field]=found[0]
    missing=[f for f in REQUIRED[kind] if f not in out]
    if missing:raise ValueError("Missing required column(s): "+", ".join(missing))
    return out

def amount(v):
    try:
        d=Decimal(str(v if v is not None else '0').replace('$','').replace(',','').strip())
        if not d.is_finite() or d<0 or d>100000000:raise ValueError()
        return int((d*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    except (ValueError,InvalidOperation):
        raise ValueError("Invalid monetary amount")

def parse(kind, values):
    def val(k):return str(values.get(k) if values.get(k) is not None else '').strip()
    if kind in ('customers','suppliers'):
        name=val('name')
        if not name:raise ValueError("Customer name required")
        result=dict(name=name,email=val('email'),phone=val('phone'),address=val('address'))
        if kind=='customers':result['notes']=val('notes')
        return result
    sku=val('sku');name=val('name')
    if not sku or not name:raise ValueError("SKU and product name required")
    if not val('price'):raise ValueError("Price required")
    try:
        q=Decimal(val('qty') or '0')
        if q!=q.to_integral_value() or q<0 or q>100000000:raise ValueError()
    except (InvalidOperation,ValueError):
        raise ValueError("Quantity must be a nonnegative whole number")
    return dict(sku=sku,name=name,barcode=val('barcode'),description=val('description'),
                price=amount(val('price')),cost=amount(val('cost') or '0'),qty=int(q),supplier=val('supplier'))

def read_excel(path,kind,sheet_name=None):
    from openpyxl import load_workbook
    if Path(path).suffix.lower() not in ('.xlsx','.xlsm'):
        raise ValueError("Select an .xlsx or .xlsm Excel workbook. Save older .xls files as .xlsx first.")
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        if sheet_name and sheet_name not in wb.sheetnames:raise ValueError("Worksheet not found")
        ws=wb[sheet_name] if sheet_name else wb.active
        rows=ws.iter_rows(values_only=True)
        try: headers=next(rows)
        except StopIteration:raise ValueError("Worksheet is empty")
        fields=mapping(headers,kind)
        parsed=[]; errors=[]
        for number,row in enumerate(rows,start=2):
            if number>MAX_ROWS+1:raise ValueError(f"Import limited to {MAX_ROWS} rows")
            if not any(v is not None and str(v).strip() for v in row):continue
            vals={k:(row[i] if i<len(row) else None) for k,i in fields.items()}
            try:parsed.append((number,parse(kind,vals)))
            except ValueError as e:errors.append((number,str(e)))
        return parsed,errors,ws.title
    finally:wb.close()

def import_rows(db,kind,rows,mode='skip'):
    """Atomic import. Customers match case-insensitive name; products match SKU.
    mode=skip or update. Inventory quantities in update mode are absolute counts.
    """
    if kind not in HEADERS or mode not in ('skip','update'):raise ValueError("Invalid import settings")
    from supplier_catalog import lookup_supplier_id, validate_supplier_rows
    if kind=='products':validate_supplier_rows(db,rows)
    created=updated=skipped=0
    # Fail fast for duplicate keys inside the workbook.
    keys=set()
    for rownum,item in rows:
        key=(item['sku'] if kind=='products' else item['name']).casefold()
        if key in keys:raise ValueError(f"Duplicate {kind[:-1]} in workbook on row {rownum}: {key}")
        keys.add(key)
    with db:
        for _,v in rows:
            if kind in ('customers','suppliers'):
                table=kind
                existing=db.execute(f"SELECT id FROM {table} WHERE name=? COLLATE NOCASE",(v['name'],)).fetchone()
                if existing:
                    if mode=='skip':skipped+=1;continue
                    db.execute(f"UPDATE {table} SET email=?,phone=?,address=?"+(",notes=?" if kind=="customers" else "")+" WHERE id=?",
                               (v['email'],v['phone'],v['address'],v['notes'],existing[0]) if kind=="customers"
                               else (v['email'],v['phone'],v['address'],existing[0]))
                    updated+=1
                else:
                    if kind=='customers':
                        db.execute("INSERT INTO customers(name,email,phone,address,notes) VALUES(?,?,?,?,?)",
                                   (v['name'],v['email'],v['phone'],v['address'],v['notes']))
                    else:
                        db.execute("INSERT INTO suppliers(name,email,phone,address) VALUES(?,?,?,?)",
                                   (v['name'],v['email'],v['phone'],v['address']))
                    created+=1
            else:
                existing=db.execute("SELECT id FROM products WHERE sku=? COLLATE NOCASE",(v['sku'],)).fetchone()
                barcode=v['barcode'] or None
                supplier_id=lookup_supplier_id(db,v.get('supplier',''))
                if existing:
                    if mode=='skip':skipped+=1;continue
                    db.execute("""UPDATE products SET name=?,description=?,price=?,cost=?,qty=?,barcode=?,
                                 supplier_id=COALESCE(?,supplier_id) WHERE id=?""",
                               (v['name'],v['description'],v['price'],v['cost'],v['qty'],barcode,supplier_id,existing[0]))
                    updated+=1
                else:
                    db.execute("""INSERT INTO products(sku,name,description,price,cost,qty,barcode,supplier_id)
                               VALUES(?,?,?,?,?,?,?,?)""",
                               (v['sku'],v['name'],v['description'],v['price'],v['cost'],v['qty'],barcode,supplier_id))
                    created+=1
    return created,updated,skipped
